"""Text extraction for stored PDF/DOCX/TXT documents.

Parsing runs on bytes that were already validated and hashed by
:mod:`research_inter_bridge.documents`. It is defensive on purpose: hostile PDFs
and DOCX zip bombs are the realistic attack surface of this phase, so page
counts, entry counts, uncompressed sizes and output length are all bounded, and
every failure produces an actionable message instead of a traceback.
"""

from __future__ import annotations

import io
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from . import documents, facts
from .documents import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE, TEXT_MEDIA_TYPE, DocumentRecord
from .errors import DocumentParsingError
from .logging import get_logger

logger = get_logger(__name__)

#: A CV is a few pages; 200 pages means the file is not what the student thinks.
MAX_PDF_PAGES = 200

#: Cap on extracted text so one hostile file cannot fill memory or the database.
MAX_EXTRACTED_CHARS = 200_000

#: Below this many characters a PDF is treated as a scan, not a text document.
MIN_PDF_TEXT_CHARS = 20


@dataclass(frozen=True, slots=True)
class ParseResult:
    """Outcome of parsing one document."""

    document: DocumentRecord
    text: str
    fact_ids: tuple[int, ...]

    @property
    def character_count(self) -> int:
        return len(self.text)


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]
    collapsed: list[str] = []
    blank = 0
    for line in lines:
        if line:
            blank = 0
            collapsed.append(line)
            continue
        blank += 1
        if blank <= 1:
            collapsed.append("")
    return "\n".join(collapsed).strip()[:MAX_EXTRACTED_CHARS]


def _pdf_text(payload: bytes) -> str:
    import pymupdf

    try:
        document = pymupdf.open(stream=payload, filetype="pdf")
    except Exception as exc:  # pymupdf raises bare exception types
        raise DocumentParsingError(f"the PDF could not be opened: {exc}") from exc

    with document:
        if document.needs_pass:
            raise DocumentParsingError(
                "the PDF is password-protected; export an unprotected copy and upload it again."
            )
        if document.page_count > MAX_PDF_PAGES:
            raise DocumentParsingError(
                f"the PDF has {document.page_count} pages; the limit is {MAX_PDF_PAGES}."
            )
        pages = [page.get_text("text") for page in document]
    text = _normalise("\n".join(pages))
    if len(text) < MIN_PDF_TEXT_CHARS:
        raise DocumentParsingError(
            "the PDF contains no selectable text (probably a scan); OCR is not available yet, "
            "so upload a text-based PDF, DOCX or TXT export."
        )
    return text


def _docx_text(payload: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(payload))
    except Exception as exc:  # python-docx raises several unrelated types
        raise DocumentParsingError(f"the DOCX could not be opened: {exc}") from exc

    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    text = _normalise("\n".join(chunks))
    if not text:
        raise DocumentParsingError(
            "the DOCX contains no text; if the CV is an image, export it as PDF or DOCX text."
        )
    return text


def _text_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = payload.decode(encoding)
        except UnicodeDecodeError:
            continue
        normalised = _normalise(text)
        if normalised:
            return normalised
    raise DocumentParsingError("the text file is empty or only whitespace.")


def extract_text(payload: bytes, media_type: str) -> str:
    """Convert validated bytes to normalised text."""

    if media_type == PDF_MEDIA_TYPE:
        return _pdf_text(payload)
    if media_type == DOCX_MEDIA_TYPE:
        return _docx_text(payload)
    if media_type == TEXT_MEDIA_TYPE:
        return _text_text(payload)
    raise DocumentParsingError(
        f"documents of type {media_type!r} cannot be parsed; "
        f"accepted: {PDF_MEDIA_TYPE}, {DOCX_MEDIA_TYPE}, {TEXT_MEDIA_TYPE}."
    )


def parse_document(
    connection: sqlite3.Connection,
    *,
    document_id: int,
    user_id: int,
    data_directory: Path | str,
    actor: str,
) -> ParseResult:
    """Parse one stored document, extract candidate facts and record the outcome.

    A parse failure is persisted (``status='failed'`` plus a short reason) and
    re-raised, so the UI can explain it while the audit trail keeps the fact that
    the file was unreadable.
    """

    record, payload = documents.read_document(
        connection,
        document_id=document_id,
        user_id=user_id,
        data_directory=data_directory,
    )

    try:
        text = extract_text(payload, record.media_type)
    except DocumentParsingError as exc:
        documents.mark_parse_failed(
            connection,
            document_id=document_id,
            user_id=user_id,
            reason=str(exc),
            actor=actor,
        )
        raise

    extracted = facts.extract_facts(text)
    fact_ids = facts.store_facts(
        connection, user_id=user_id, document_id=document_id, facts=extracted
    )
    documents.mark_parsed(connection, document_id=document_id, user_id=user_id, actor=actor)

    logger.info(
        "document parsed document_id=%s user_id=%s chars=%s facts=%s",
        document_id,
        user_id,
        len(text),
        len(fact_ids),
    )
    parsed = documents.get_document(connection, document_id=document_id, user_id=user_id)
    return ParseResult(document=parsed, text=text, fact_ids=fact_ids)
