"""Phase 2b — document validation, storage, isolation and parsing."""

from __future__ import annotations

import io
import os
import sqlite3
from pathlib import Path

import pymupdf
import pytest
from docx import Document as WordDocument

from research_inter_bridge import db, documents, parsing
from research_inter_bridge.documents import (
    DOCX_MEDIA_TYPE,
    PDF_MEDIA_TYPE,
    TEXT_MEDIA_TYPE,
    DocumentRecord,
)
from research_inter_bridge.errors import (
    DocumentError,
    DocumentNotFoundError,
    DocumentParsingError,
    DocumentTooLargeError,
    UnsupportedDocumentError,
)
from research_inter_bridge.profiles import (
    confirmed_facts,
    create_user,
    delete_user,
    list_profile_facts,
)


@pytest.fixture
def connection(tmp_path: Path):
    with db.session(tmp_path / "app.db") as conn:
        db.apply_migrations(conn)
        yield conn


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    return tmp_path / "data"


def pdf_bytes(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text, fontsize=11)
    payload = document.tobytes()
    document.close()
    return payload


def blank_pdf_bytes() -> bytes:
    """A PDF with no selectable text, like a scan."""

    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(pymupdf.Rect(72, 72, 400, 400), color=(0, 0, 0), width=2)
    payload = document.tobytes()
    document.close()
    return payload


def docx_bytes(*paragraphs: str) -> bytes:
    document = WordDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(
    connection: sqlite3.Connection,
    data_dir: Path,
    user_id: int,
    payload: bytes,
    name: str = "cv.pdf",
    **kwargs,
) -> DocumentRecord:
    return documents.save_document(
        connection,
        user_id=user_id,
        display_name=name,
        payload=payload,
        data_directory=data_dir,
        actor=f"student:{user_id}",
        **kwargs,
    )


# ------------------------------------------------------------------- validation


def test_pdf_is_stored_under_the_data_directory_with_a_hash_name(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    payload = pdf_bytes("Asha Verma")
    record = _upload(connection, data_dir, user_id, payload)

    assert record.media_type == PDF_MEDIA_TYPE
    assert record.extension == ".pdf"
    assert record.status == "stored"
    assert record.display_name == "cv.pdf"
    assert record.storage_path.startswith(f"user_{user_id}/")
    assert record.storage_path.endswith(".pdf")
    assert record.sha256[:16] in record.storage_path

    stored = data_dir / "documents" / record.storage_path
    assert stored.is_file()
    assert stored.read_bytes() == payload


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("../../etc/passwd.pdf", "passwd.pdf"),
        (r"..\..\windows\system32\cv.pdf", "cv.pdf"),
        ("  my  cv.PDF  ", "my cv.PDF"),
        ("...hidden.docx", "hidden.docx"),
    ],
)
def test_client_file_name_cannot_escape_the_storage_root(
    connection: sqlite3.Connection, data_dir: Path, name: str, expected: str
) -> None:
    user_id = create_user(connection)
    payload = docx_bytes("Asha Verma") if name.lower().endswith(".docx") else pdf_bytes("Asha")
    record = _upload(connection, data_dir, user_id, payload, name=name)

    assert record.display_name == expected
    assert ".." not in record.storage_path
    stored = (data_dir / "documents" / record.storage_path).resolve()
    assert stored.is_relative_to((data_dir / "documents").resolve())


def test_wrong_extension_lists_the_allowed_set(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    with pytest.raises(UnsupportedDocumentError, match=r"\.pdf"):
        _upload(connection, data_dir, user_id, b"%PDF-1.7 ...", name="cv.exe")


def test_executable_masquerading_as_pdf_is_rejected(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    with pytest.raises(UnsupportedDocumentError, match="content"):
        _upload(connection, data_dir, user_id, b"MZ\x90\x00\x03\x00\x00\x00", name="cv.pdf")


def test_zip_that_is_not_a_word_document_is_rejected(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hello")
    user_id = create_user(connection)
    with pytest.raises(UnsupportedDocumentError):
        _upload(connection, data_dir, user_id, buffer.getvalue(), name="cv.docx")


def test_oversized_upload_is_rejected_before_parsing(
    connection: sqlite3.Connection, data_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(documents, "MAX_DOCUMENT_BYTES", 64)
    user_id = create_user(connection)
    with pytest.raises(DocumentTooLargeError, match="limit is 64 bytes"):
        _upload(connection, data_dir, user_id, pdf_bytes("x" * 200))


def test_empty_file_is_rejected(connection: sqlite3.Connection, data_dir: Path) -> None:
    user_id = create_user(connection)
    with pytest.raises(UnsupportedDocumentError, match="empty"):
        _upload(connection, data_dir, user_id, b"", name="cv.txt")


def test_declared_media_type_must_match_content(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    with pytest.raises(UnsupportedDocumentError, match="declares"):
        _upload(
            connection,
            data_dir,
            user_id,
            pdf_bytes("Asha"),
            declared_media_type="text/plain",
        )
    # Browsers often send the generic type; that must still work.
    record = _upload(
        connection,
        data_dir,
        user_id,
        pdf_bytes("Asha"),
        declared_media_type="application/octet-stream",
    )
    assert record.media_type == PDF_MEDIA_TYPE


def test_upload_for_unknown_user_is_rejected(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    with pytest.raises(DocumentError, match="create the user first"):
        _upload(connection, data_dir, 42, pdf_bytes("Asha"))


def test_unknown_kind_is_rejected(connection: sqlite3.Connection, data_dir: Path) -> None:
    user_id = create_user(connection)
    with pytest.raises(DocumentError, match="kind"):
        _upload(connection, data_dir, user_id, pdf_bytes("Asha"), kind="passport")


# -------------------------------------------------------------------- storage


def test_reuploading_identical_bytes_is_idempotent(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    payload = pdf_bytes("Asha Verma")
    first = _upload(connection, data_dir, user_id, payload)
    second = _upload(connection, data_dir, user_id, payload, name="renamed.pdf")

    assert first.id == second.id
    assert first.display_name == "cv.pdf"
    assert len(documents.list_documents(connection, user_id)) == 1


def test_two_users_may_store_the_same_file_independently(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    first = create_user(connection)
    second = create_user(connection)
    payload = pdf_bytes("Shared template")
    a = _upload(connection, data_dir, first, payload)
    b = _upload(connection, data_dir, second, payload)

    assert a.id != b.id
    assert a.storage_path != b.storage_path


def test_documents_are_isolated_per_user(connection: sqlite3.Connection, data_dir: Path) -> None:
    owner = create_user(connection)
    stranger = create_user(connection)
    record = _upload(connection, data_dir, owner, pdf_bytes("Asha"))

    assert documents.list_documents(connection, stranger) == ()
    with pytest.raises(DocumentNotFoundError):
        documents.get_document(connection, document_id=record.id, user_id=stranger)
    with pytest.raises(DocumentNotFoundError):
        documents.read_document(
            connection, document_id=record.id, user_id=stranger, data_directory=data_dir
        )
    assert (
        documents.delete_document(
            connection,
            document_id=record.id,
            user_id=stranger,
            data_directory=data_dir,
            actor="student:2",
        )
        is False
    )


def test_delete_removes_row_and_file(connection: sqlite3.Connection, data_dir: Path) -> None:
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, pdf_bytes("Asha"))
    path = data_dir / "documents" / record.storage_path
    assert path.is_file()

    assert (
        documents.delete_document(
            connection,
            document_id=record.id,
            user_id=user_id,
            data_directory=data_dir,
            actor="student:1",
        )
        is True
    )
    assert not path.exists()
    assert documents.list_documents(connection, user_id) == ()


def test_purge_user_documents_removes_files_and_rows(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    first = _upload(connection, data_dir, user_id, pdf_bytes("one"))
    second = _upload(connection, data_dir, user_id, docx_bytes("two"), name="cv.docx")
    paths = [data_dir / "documents" / record.storage_path for record in (first, second)]

    removed = documents.purge_user_documents(
        connection, user_id=user_id, data_directory=data_dir, actor="student:1"
    )

    assert removed == 2
    assert not any(path.exists() for path in paths)
    assert documents.list_documents(connection, user_id) == ()


def test_tampered_file_fails_its_integrity_check(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, pdf_bytes("Asha"))
    (data_dir / "documents" / record.storage_path).write_bytes(pdf_bytes("Mallory"))

    with pytest.raises(DocumentError, match="integrity"):
        documents.read_document(
            connection, document_id=record.id, user_id=user_id, data_directory=data_dir
        )


def test_missing_file_gives_an_actionable_error(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, pdf_bytes("Asha"))
    (data_dir / "documents" / record.storage_path).unlink()

    with pytest.raises(DocumentError, match="upload it again"):
        documents.read_document(
            connection, document_id=record.id, user_id=user_id, data_directory=data_dir
        )


def test_stored_file_is_not_world_readable(connection: sqlite3.Connection, data_dir: Path) -> None:
    if os.name == "nt":  # POSIX mode bits are not meaningful on Windows
        pytest.skip("POSIX permission bits are not enforced on Windows")
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, pdf_bytes("Asha"))
    mode = (data_dir / "documents" / record.storage_path).stat().st_mode & 0o777
    assert mode == 0o600


def test_audit_events_record_ids_and_hashes_but_never_content(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, pdf_bytes("Asha Verma secret"))

    events = db.list_audit_events(connection)
    uploaded = [event for event in events if event.action == "document.uploaded"]
    assert len(uploaded) == 1
    assert uploaded[0].target_id == str(record.id)
    assert uploaded[0].metadata["sha256"] == record.sha256
    assert "secret" not in repr(uploaded[0].metadata)


# -------------------------------------------------------------------- parsing


def test_pdf_text_is_extracted_and_facts_are_pending(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    record = _upload(
        connection,
        data_dir,
        user_id,
        pdf_bytes("Asha Verma asha@example.edu CGPA: 8.9/10 Expected graduation 2027"),
    )

    result = parsing.parse_document(
        connection,
        document_id=record.id,
        user_id=user_id,
        data_directory=data_dir,
        actor="student:1",
    )

    assert "Asha Verma" in result.text
    assert result.document.status == "parsed"
    assert result.fact_ids
    fields = {fact["field"] for fact in list_profile_facts(connection, user_id)}
    assert {"contact_email", "gpa", "expected_graduation_year"} <= fields
    assert confirmed_facts(connection, user_id) == ()


def test_docx_text_including_tables_is_extracted(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    document = WordDocument()
    document.add_paragraph("Asha Verma")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "CGPA"
    table.cell(0, 1).text = "9.1/10"
    buffer = io.BytesIO()
    document.save(buffer)

    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, buffer.getvalue(), name="cv.docx")
    result = parsing.parse_document(
        connection,
        document_id=record.id,
        user_id=user_id,
        data_directory=data_dir,
        actor="student:1",
    )

    assert "Asha Verma" in result.text
    assert "9.1/10" in result.text
    assert result.document.media_type == DOCX_MEDIA_TYPE


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_text_file_encodings_are_detected(
    connection: sqlite3.Connection, data_dir: Path, encoding: str
) -> None:
    user_id = create_user(connection)
    payload = "Asha Verma\nGPA 3.7/4\n".encode(encoding)
    record = _upload(connection, data_dir, user_id, payload, name="cv.txt")

    result = parsing.parse_document(
        connection,
        document_id=record.id,
        user_id=user_id,
        data_directory=data_dir,
        actor="student:1",
    )

    assert "Asha Verma" in result.text
    assert result.document.media_type == TEXT_MEDIA_TYPE


def test_scanned_pdf_is_marked_failed_with_ocr_explanation(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, blank_pdf_bytes())

    with pytest.raises(DocumentParsingError, match="OCR"):
        parsing.parse_document(
            connection,
            document_id=record.id,
            user_id=user_id,
            data_directory=data_dir,
            actor="student:1",
        )

    stored = documents.get_document(connection, document_id=record.id, user_id=user_id)
    assert stored.status == "failed"
    assert "OCR" in (stored.parse_error or "")
    actions = [event.action for event in db.list_audit_events(connection)]
    assert "document.parse_failed" in actions


def test_password_protected_pdf_is_rejected(connection: sqlite3.Connection, data_dir: Path) -> None:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Asha Verma secret")
    payload = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="student"
    )
    document.close()

    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, payload)

    with pytest.raises(DocumentParsingError, match="password"):
        parsing.parse_document(
            connection,
            document_id=record.id,
            user_id=user_id,
            data_directory=data_dir,
            actor="student:1",
        )


def test_reparsing_is_idempotent_for_facts(connection: sqlite3.Connection, data_dir: Path) -> None:
    user_id = create_user(connection)
    record = _upload(connection, data_dir, user_id, pdf_bytes("Asha asha@example.edu CGPA 8/10"))

    first = parsing.parse_document(
        connection,
        document_id=record.id,
        user_id=user_id,
        data_directory=data_dir,
        actor="student:1",
    )
    second = parsing.parse_document(
        connection,
        document_id=record.id,
        user_id=user_id,
        data_directory=data_dir,
        actor="student:1",
    )

    assert first.fact_ids == second.fact_ids
    assert len(list_profile_facts(connection, user_id)) == len(first.fact_ids)


def test_parsing_cannot_touch_another_users_document(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    owner = create_user(connection)
    stranger = create_user(connection)
    record = _upload(connection, data_dir, owner, pdf_bytes("Asha"))

    with pytest.raises(DocumentNotFoundError):
        parsing.parse_document(
            connection,
            document_id=record.id,
            user_id=stranger,
            data_directory=data_dir,
            actor="student:2",
        )


def test_deleting_a_user_cascades_document_rows(
    connection: sqlite3.Connection, data_dir: Path
) -> None:
    user_id = create_user(connection)
    _upload(connection, data_dir, user_id, pdf_bytes("Asha"))

    delete_user(connection, user_id, actor="student:1")

    remaining = connection.execute("SELECT COUNT(*) AS c FROM student_documents").fetchone()["c"]
    assert remaining == 0
